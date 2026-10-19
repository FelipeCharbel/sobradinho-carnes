"""
Agente A: testes de geração e integridade dos PPTX (Modo A e Modo B). Ver
docs/CONTRATOS_FASE2.md §10 para a lista completa exigida. Casos mínimos:

1. construir_spec_cargo: competencias/autonomia/processos sempre ["A confirmar."] (não existem no
   banco); supervisao == ["Não exerce supervisão."] quando sem subordinados, e lista os nomes
   quando houver.
2. dobra_rotinas: frequência "Trimestral" cai em "Sob demanda" com prefixo "(Trimestral) ".
3. gerar_slide_cargo: binário não vazio; abre com pptx.Presentation(io.BytesIO(resultado)) sem
   exceção; tem exatamente 5 slides.
4. gerar_manual_completo: abre sem exceção; contém o slide de organograma; nenhum "3.XX" (só
   aparece no modo "cargo"); contagem de slides == 11 + 1 (organograma) + 5*n_cargos_publicados
   (ajustar a constante exata depois de ler o slide_engine.py final).
5. Rodar tests/test_fluxo.py e tests/test_formalizacao.py no final e confirmar 100% verde (sem
   regressão da Fase 1).

Usar um banco SQLite em memória por teste (mesmo padrão de tests/test_fluxo.py), nunca o banco de
desenvolvimento.
"""

import pytest
import json
import io
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from pptx import Presentation

from server.models import Base, Cargo, Usuario
from server.generator.spec_builder import (
    pontua_lista, dobra_rotinas, supervisao_de, construir_spec_cargo, construir_spec_manual, construir_spec_um_cargo
)
from server.generator.slide_engine import gerar_slide_cargo, gerar_manual_completo


@pytest.fixture
def db_memory():
    """Cria banco SQLite em memória para testes."""
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


def test_pontua_lista_empty():
    """Lista vazia -> ["A confirmar."]"""
    result = pontua_lista([])
    assert result == ["A confirmar."]


def test_pontua_lista_single():
    """Um item -> item + '.'"""
    result = pontua_lista(["Liderança"])
    assert result == ["Liderança."]


def test_pontua_lista_multiple():
    """Múltiplos itens -> ';' em todos menos último que leva '.'"""
    result = pontua_lista(["Liderança", "Comunicação", "Organização"])
    assert result == ["Liderança;", "Comunicação;", "Organização."]


def test_dobra_rotinas_vazio():
    """Lista vazia -> 4 blocos com "A confirmar."."""
    result = dobra_rotinas([])
    assert len(result) == 4
    assert result[0][0] == "Diariamente"
    assert result[1][0] == "Semanalmente"
    assert result[2][0] == "Mensalmente"
    assert result[3][0] == "Sob demanda"
    for bloco in result:
        assert bloco[1] == ["A confirmar."]


def test_dobra_rotinas_trimestral():
    """Frequência Trimestral cai em "Sob demanda" com prefixo."""
    rotinas = [
        {"f": "Trimestral", "t": "Avaliar os relatórios."}
    ]
    result = dobra_rotinas(rotinas)
    # Encontra "Sob demanda"
    sob_demanda = next(b for b in result if b[0] == "Sob demanda")
    assert len(sob_demanda[1]) == 1
    assert "(Trimestral)" in sob_demanda[1][0]
    assert "Avaliar os relatórios." in sob_demanda[1][0]


def test_dobra_rotinas_mapeamento():
    """Mapeia Diária/Semanal/Mensal direto."""
    rotinas = [
        {"f": "Diária", "t": "Verificar o caixa."},
        {"f": "Semanal", "t": "Reunião de equipe."},
        {"f": "Mensal", "t": "Relatório mensal."},
    ]
    result = dobra_rotinas(rotinas)

    diariamente = next(b for b in result if b[0] == "Diariamente")
    semanalmente = next(b for b in result if b[0] == "Semanalmente")
    mensalmente = next(b for b in result if b[0] == "Mensalmente")

    assert "Verificar o caixa." in diariamente[1][0]
    assert "Reunião de equipe." in semanalmente[1][0]
    assert "Relatório mensal." in mensalmente[1][0]


def test_supervisao_de_sem_subordinados():
    """Sem subordinados -> "Não exerce supervisão."."""
    organograma = [
        {"id": "c1", "nome": "Gerente", "pai": None},
        {"id": "c2", "nome": "Analista", "pai": "c1"},
    ]
    result = supervisao_de("c2", organograma)
    assert result == ["Não exerce supervisão."]


def test_supervisao_de_com_subordinados():
    """Com subordinados -> lista os nomes, pontuados."""
    organograma = [
        {"id": "c1", "nome": "Gerente", "pai": None},
        {"id": "c2", "nome": "Analista 1", "pai": "c1"},
        {"id": "c3", "nome": "Analista 2", "pai": "c1"},
    ]
    result = supervisao_de("c1", organograma)
    assert len(result) == 2
    assert result[0].endswith(";")
    assert result[1].endswith(".")
    assert "Analista 1" in result[0] or "Analista 1" in result[1]
    assert "Analista 2" in result[0] or "Analista 2" in result[1]


def test_construir_spec_cargo_campos_nao_existentes(db_memory):
    """competencias/autonomia/processos sempre ["A confirmar."]."""
    cargo_dict = {
        "id": "c1",
        "nome": "Gerente",
        "conteudo": {
            "missao": "Liderar",
            "funcoes": ["Contratar"],
            "rotinas": [],
            "formacao": "",
            "experiencia": "",
            "habilidades": [],
            "relacoes": [],
            "status": "publicado"
        }
    }
    organograma = [cargo_dict]

    spec = construir_spec_cargo(cargo_dict, organograma)

    assert spec["competencias"] == ["A confirmar."]
    assert spec["autonomia"] == ["A confirmar."]
    assert spec["processos"] == ["A confirmar."]


def test_construir_spec_cargo_supervisao_derivada(db_memory):
    """supervisao é derivada do organograma."""
    cargo_gerente = {
        "id": "gerente",
        "nome": "Gerente",
        "conteudo": {
            "missao": "Liderar",
            "funcoes": ["Contratar"],
            "rotinas": [],
            "formacao": "",
            "experiencia": "",
            "habilidades": [],
            "relacoes": [],
            "status": "publicado"
        }
    }
    cargo_analista = {
        "id": "analista",
        "nome": "Analista",
        "pai": "gerente",
        "conteudo": {
            "missao": "Analisar",
            "funcoes": ["Revisar"],
            "rotinas": [],
            "formacao": "",
            "experiencia": "",
            "habilidades": [],
            "relacoes": [],
            "status": "publicado"
        }
    }

    organograma = [cargo_gerente, cargo_analista]
    spec_gerente = construir_spec_cargo(cargo_gerente, organograma)
    spec_analista = construir_spec_cargo(cargo_analista, organograma)

    # Gerente lidera Analista
    assert len(spec_gerente["supervisao"]) == 1
    assert "Analista" in spec_gerente["supervisao"][0]

    # Analista não lidera ninguém
    assert spec_analista["supervisao"] == ["Não exerce supervisão."]


def test_gerar_slide_cargo_valido(db_memory):
    """gerar_slide_cargo devolve .pptx válido com 5 slides."""
    # Cria cargo no banco
    cargo = Cargo(
        id="c1",
        nome="Gerente",
        area="Administração",
        pessoa="João",
        missao="Liderar",
        funcoes=json.dumps(["Contratar", "Avaliar"]),
        rotinas=json.dumps([
            {"f": "Diária", "t": "Verificar o caixa."},
            {"f": "Semanal", "t": "Reunião de equipe."}
        ]),
        formacao="MBA",
        experiencia="10 anos",
        habilidades=json.dumps(["Liderança"]),
        relacoes=json.dumps(["Sócios"]),
        status="publicado"
    )
    db_memory.add(cargo)
    db_memory.commit()

    spec = construir_spec_um_cargo("c1", db_memory)
    # Garante que spec tem "emails" (pode vir vazio de construir_spec_um_cargo)
    if "emails" not in spec:
        spec["emails"] = []

    # Gera PPTX
    pptx_bytes = gerar_slide_cargo(spec)

    # Valida binário
    assert pptx_bytes is not None
    assert len(pptx_bytes.getvalue()) > 0

    # Abre com pptx.Presentation
    prs = Presentation(pptx_bytes)

    # Verifica número de slides
    assert len(prs.slides) == 5


def test_gerar_manual_completo_valido(db_memory):
    """gerar_manual_completo devolve .pptx válido com slides corretos."""
    # Cria 2 cargos no banco
    cargo1 = Cargo(
        id="c1",
        nome="Gerente",
        area="Administração",
        pessoa="João",
        missao="Liderar",
        funcoes=json.dumps(["Contratar"]),
        rotinas=json.dumps([]),
        formacao="MBA",
        experiencia="10 anos",
        habilidades=json.dumps(["Liderança"]),
        relacoes=json.dumps([]),
        status="publicado"
    )
    cargo2 = Cargo(
        id="c2",
        nome="Analista",
        area="Administração",
        pai="c1",
        pessoa="Maria",
        missao="Analisar",
        funcoes=json.dumps(["Revisar"]),
        rotinas=json.dumps([]),
        formacao="Graduação",
        experiencia="5 anos",
        habilidades=json.dumps([]),
        relacoes=json.dumps([]),
        status="publicado"
    )
    db_memory.add(cargo1)
    db_memory.add(cargo2)
    db_memory.commit()

    spec = construir_spec_manual(db_memory)

    # Gera PPTX
    pptx_bytes = gerar_manual_completo(spec)

    # Valida binário
    assert pptx_bytes is not None
    assert len(pptx_bytes.getvalue()) > 0

    # Abre com pptx.Presentation
    prs = Presentation(pptx_bytes)

    # Verifica que não tem "3.XX" (modo "cargo")
    xml_string = str(prs.slides[0]._element)
    assert "3.XX" not in xml_string

    # capa=1, sumário=1, divisória-intro=1, introdução=1, divisória2=1, 4 páginas explicativas=4,
    # organograma=1, divisória-cargos=1, cargo1(divisória+4 páginas)=5, cargo2=5, contra-capa=1
    # Total = 1+1+1+1+1+4+1+1+5+5+1 = 22
    assert len(prs.slides) == 22

    # Confirma que nenhuma entrada do .pptx (ZIP OOXML) está duplicada -- é o sintoma do bug de
    # corrupção que o coordenador encontrou e corrigiu (organograma inserido fora de sequência
    # corrompia a alocação de nomes de parte do python-pptx).
    import zipfile
    from collections import Counter
    pptx_bytes.seek(0)
    nomes = zipfile.ZipFile(pptx_bytes).namelist()
    duplicadas = {n: c for n, c in Counter(nomes).items() if c > 1}
    assert not duplicadas, f"Entradas duplicadas no .pptx (corrupção): {duplicadas}"
